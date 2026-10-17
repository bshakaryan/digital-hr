"""Structural QA for the five DOCX files and independent evaluation CSV."""
from __future__ import annotations

import csv
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from build_corpus import DOCS

ROOT = Path(__file__).resolve().parent
required = ("Document ID:", "Version / created:", "Category:", "Source classification:",
            "Intended audience:", "Access scope:", "DEMO / SYNTHETIC")
supports = set()
names = []

for item in DOCS:
    path = ROOT / f"{item['id']}.docx"
    names.append(path.name)
    with ZipFile(path) as package:
        assert package.testzip() is None
        assert "word/document.xml" in package.namelist()
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    assert paragraphs[0] == item["id"]
    assert item["title"] in paragraphs[1]
    for marker in required:
        assert marker in text, (item["id"], marker)
    assert f"Access scope: {item['scope']}" in text
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    for sid, _, bodies in item["sections"]:
        assert len(bodies) > 0
        assert any(re.match(rf"^{re.escape(item['id'])} {re.escape(sid)}\s", h) for h in headings), (item["id"], sid)
        supports.add(f"{item['id']} {sid}")
    assert len(headings) == len(item["sections"]) + 1
    assert len(text) > 2500, item["id"]
    print(f"{path.name}: {len(item['sections'])} numbered sections, {len(text)} extracted characters, ZIP and python-docx OK")

with ZipFile(ROOT / "ai-jana-corpus.zip") as archive:
    assert sorted(archive.namelist()) == sorted(names)
    for name in names:
        assert archive.read(name) == (ROOT / name).read_bytes()

with (ROOT / "rag_evaluation.csv").open(encoding="utf-8", newline="") as handle:
    questions = list(csv.DictReader(handle))
assert 25 <= len(questions) <= 32
assert len({row["id"] for row in questions}) == len(questions)
for row in questions:
    if row["expected_outcome"] == "supported":
        assert row["expected_support"]
        for ref in row["expected_support"].split("|"):
            assert ref in supports, (row["id"], ref)
    else:
        assert not row["expected_support"], row["id"]

with (ROOT / "traceability.csv").open(encoding="utf-8", newline="") as handle:
    claims = list(csv.DictReader(handle))
assert {row["classification"] for row in claims} >= {
    "verified_app", "general_education", "synthetic_demo", "organizer_requirement", "evidence_gap", "unsupported"}
print(f"ZIP: exactly 5 matching DOCX files; evaluation: {len(questions)} questions; matrix: {len(claims)} claims")
