from __future__ import annotations

import hashlib
import io
from pathlib import Path

import pytest
from docx import Document

from app.modules.knowledge_ingestion import extract_docx, extract_pdf


CORPUS = Path(__file__).resolve().parents[2] / "knowledge_base_docs"


def docx_bytes(text: str, heading="01. Материал") -> bytes:
    document = Document()
    document.add_paragraph("DEMO / SYNTHETIC — учебный пример.")
    document.add_heading(heading, level=1)
    document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def upload(api, data: bytes, *, filename="test.docx", title="Учебный тест", scope="all", actor="hr"):
    return api.client.post("/api/knowledge-files", headers=api.headers(actor),
        data={"title": title, "version": "test-1", "language": "kk" if "қазақ" in title else "ru",
              "classification": "synthetic_demo", "access_scope": scope,
              "source_note": "Только тест", "document_code": "TEST-01"},
        files={"file": (filename, data)})


def test_docx_upload_hash_dedup_original_download_and_kazakh(api):
    text = "Қазақ тіліндегі қоймаға келіп түскен сәлемдеме тіркелді. Кириллический текст сохраняется без потерь."
    data = docx_bytes(text)
    before = api.db.snapshot("knowledge_documents")
    result = upload(api, data, title="қазақ тіліндегі тест")
    assert result.status_code == 200, result.text
    file = result.json()
    assert file["sha256"] == hashlib.sha256(data).hexdigest()
    assert file["chunk_count"] >= 1 and not file["duplicate"]
    chunk = api.db.snapshot("knowledge_documents")["knowledge_documents"]
    assert len(chunk) > len(before["knowledge_documents"])
    assert "Қазақ" in " ".join(part.content for part in extract_docx(data))
    source_id = next(row.id for row in chunk if row.file_id == file["id"] and row.section == "01. Материал")
    source = api.get(f"/api/knowledge/{source_id}").json()
    assert source["source_filename"] == "test.docx"
    assert source["section"] == "01. Материал"
    assert source["classification"] == "synthetic_demo"
    download = api.get(f"/api/knowledge-files/{file['id']}/download")
    assert download.content == data
    repeat = upload(api, data, title="другое название", scope="hr_admin")
    assert repeat.status_code == 200 and repeat.json()["duplicate"]
    assert repeat.json()["id"] == file["id"]
    assert api.db.snapshot("knowledge_documents")["knowledge_documents"] == chunk


def test_only_hr_uploads_and_restricted_file_never_reaches_employee(api):
    data = docx_bytes("Секретный учебный термин: көкжиек қоймасы. Он есть только в этом тестовом файле.")
    assert upload(api, data, actor="employee").status_code == 403
    result = upload(api, data, scope="hr_admin")
    assert result.status_code == 200, result.text
    file_id = result.json()["id"]
    assert api.get("/api/knowledge-files", "employee").status_code == 403
    assert api.get(f"/api/knowledge-files/{file_id}", "employee").status_code == 404
    assert api.get(f"/api/knowledge-files/{file_id}/download", "employee").status_code == 404
    source_id = next(row.id for row in api.db.snapshot("knowledge_documents")["knowledge_documents"] if row.file_id == file_id)
    assert api.get(f"/api/knowledge/{source_id}", "employee").status_code == 404
    answer = api.post("/api/assistant/query", {"question": "көкжиек қоймасы"}).json()
    assert all(source.get("file_id") != file_id for source in answer["sources"])
    hr_answer = api.post("/api/assistant/query", {"question": "көкжиек қоймасы"}, "hr").json()
    assert any(source.get("file_id") == file_id for source in hr_answer["sources"])


def test_course_question_uses_only_attached_authorized_files(api):
    first = upload(api, docx_bytes("Белгіленген жәшік сұрыптау аймағына көшірілді. Көкшетау коды оқылды."),
                   title="қазақ сұрыптау").json()
    second_response = upload(api, docx_bytes("Айрықша санаттағы тауар тіркелді. Жарқырау коды тек осы құжатта. Бұл оқу мысалындағы код басқа құжатқа немесе курсқа қатысы жоқ."),
                             title="қазақ басқа")
    assert second_response.status_code == 200, second_response.text
    second = second_response.json()
    assert api.post("/api/courses/1/materials", {"file_id": first["id"]}, "employee").status_code == 403
    attached = api.post("/api/courses/1/materials", {"file_id": first["id"]}, "hr")
    assert attached.status_code == 200, attached.text
    before = api.db.snapshot("course_materials", "audit_events")
    assert api.post("/api/courses/1/materials", {"file_id": first["id"]}, "hr").status_code == 200
    assert api.db.snapshot("course_materials", "audit_events") == before
    materials = api.get("/api/courses/1/materials").json()
    assert [item["id"] for item in materials] == [first["id"]]
    answer = api.post("/api/courses/1/ask", {"question": "Көкшетау коды"}).json()
    assert answer["sources"] and all(source["file_id"] == first["id"] for source in answer["sources"])
    unrelated = api.post("/api/courses/1/ask", {"question": "Жарқырау коды"}).json()
    assert all(source.get("file_id") != second["id"] for source in unrelated["sources"])
    assert api.get("/api/courses/99999/materials").status_code == 404


def test_pdf_cyrillic_pages_and_citations(api):
    path = CORPUS / "HR-01.pdf"
    if not path.exists():
        pytest.skip("Optional provided corpus is not mounted")
    data = path.read_bytes()
    chunks = extract_pdf(data)
    assert {chunk.page_number for chunk in chunks} == {1, 2, 3, 4}
    assert any("отпуск" in chunk.content.lower() for chunk in chunks)
    result = upload(api, data, filename="HR-01.pdf", title="Подача и согласование отпуска")
    assert result.status_code == 200, result.text
    file_id = result.json()["id"]
    rows = [row for row in api.db.snapshot("knowledge_documents")["knowledge_documents"] if row.file_id == file_id]
    assert rows and all(row.page_number is not None for row in rows)
    source = api.get(f"/api/knowledge/{rows[0].id}").json()
    assert source["source_filename"] == "HR-01.pdf" and source["page_number"] == rows[0].page_number
