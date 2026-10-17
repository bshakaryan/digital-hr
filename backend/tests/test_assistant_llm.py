"""Mocked provider and API tests; no external LLM calls or persistent data."""

import json
import io

import httpx

from app.config import Settings
from app.modules.assistant_llm import generate_answer


def settings(**changes):
    values = {"database_url": "sqlite:///test-only.sqlite3",
              "jwt_secret": "test-only-jwt-secret-never-use-outside-pytest-2026",
              "llm_provider": "openai_compatible", "llm_api_key": "test-provider-key",
              "openai_api_key": None}
    values.update(changes)
    return Settings(_env_file=None, **values)


def completion(text, reason="stop"):
    return {"choices": [{"finish_reason": reason, "message": {"content": text}}]}


def test_provider_receives_bounded_evidence_and_returns_safe_wording():
    requests = []

    def respond(request):
        requests.append(request)
        return httpx.Response(200, json=completion("Заявку на отпуск можно отправить через сервис заявок."))

    passage = "Заявку на отпуск сотрудник отправляет через сервис заявок."
    answer = generate_answer("Как отправить отпускную заявку?", passage, settings(),
                             transport=httpx.MockTransport(respond))
    assert answer == "Заявку на отпуск можно отправить через сервис заявок."
    request = requests[0]
    assert request.url == "https://api.openai.com/v1/chat/completions"
    assert request.headers["Authorization"] == "Bearer test-provider-key"
    body = json.loads(request.content)
    assert body["max_completion_tokens"] == 200
    assert body["messages"][1]["content"].count(passage) == 1
    assert "tools" not in body


def test_provider_errors_and_unsafe_outputs_use_fallback():
    passage = "Заявку на отпуск сотрудник отправляет через сервис заявок."
    for response in (
        httpx.Response(503),
        httpx.Response(200, json=completion("Посмотрите https://example.com.")),
        httpx.Response(200, json=completion("Посмотрите example.com.")),
        httpx.Response(200, json=completion("Посмотрите example.xyz.")),
        httpx.Response(200, json=completion("Отпуск длится 999 дней.")),
        httpx.Response(200, json=completion("Заявка одобрена.")),
        httpx.Response(200, json=completion("Заявку на отпуск можно отправить", "length")),
    ):
        result = generate_answer("Как отправить отпускную заявку?", passage, settings(),
                                 transport=httpx.MockTransport(lambda _: response))
        assert result is None
    assert generate_answer("Как отправить отпускную заявку?", passage,
                           settings(llm_api_key=None)) is None
    assert generate_answer("Как отправить отпускную заявку?", passage,
                           settings(llm_api_key=None, openai_api_key="test-provider-key"),
                           transport=httpx.MockTransport(lambda _: httpx.Response(200, json=completion("NO_EVIDENCE")))) == "NO_EVIDENCE"
    assert generate_answer("Как отправить отпускную заявку?", passage, settings(),
                           transport=httpx.MockTransport(lambda _: httpx.Response(200, json=completion("NO_EVIDENCE")))) == "NO_EVIDENCE"


def test_api_keeps_server_citation_and_extracts_on_provider_failure(api, monkeypatch):
    from app.modules import assistant

    seen = []

    def generated(question, passage, settings):
        seen.append((question, passage))
        return "Заявку на отпуск можно подать через форму в приложении."

    monkeypatch.setattr(assistant, "generate_answer", generated)
    result = api.post("/api/assistant/query", {"question": "Как подать заявку на отпуск?"}).json()
    assert result["answer"].endswith("Заявку на отпуск можно подать через форму в приложении. [1]")
    assert result["mode"] == "extractive_demo"
    assert result["sources"] == [api.get(f"/api/knowledge/{result['sources'][0]['document_id']}").json()]
    assert len(seen) == 1

    monkeypatch.setattr(assistant, "generate_answer", lambda *args: None)
    fallback = api.post("/api/assistant/query", {"question": "Как подать заявку на отпуск?"}).json()
    assert fallback["sources"] == result["sources"]
    assert fallback["answer"].endswith(" [1]")
    assert fallback["answer"] != result["answer"]


def test_no_evidence_or_access_never_reaches_provider(api, monkeypatch):
    from app.modules import assistant

    calls = []
    monkeypatch.setattr(assistant, "generate_answer", lambda *args: calls.append(args))
    assert api.post("/api/assistant/query", {"question": "Какова орбита Нептуна?"}).json()["sources"] == []
    assert api.post("/api/assistant/query", {"question": "аудит журнал"}).json()["sources"] == []
    assert api.post("/api/courses/1/ask", {"question": "Какова орбита Нептуна?"}).json()["sources"] == []
    assert api.post("/api/courses/1/ask", {"question": "Где увидеть решение по заявке?"}).json()["sources"] == []
    assert calls == []


def test_model_declared_insufficient_evidence_refuses_with_no_source(api, monkeypatch):
    from app.modules import assistant

    monkeypatch.setattr(assistant, "generate_answer", lambda *args: "NO_EVIDENCE")
    result = api.post("/api/assistant/query", {"question": "Как подать заявку на отпуск?"}).json()
    assert result["escalate_to_hr"] is True
    assert result["sources"] == []


def test_course_llm_context_contains_only_attached_authorized_material(api, monkeypatch):
    from docx import Document
    from app.modules import assistant

    def upload(text, name):
        document = Document()
        document.add_paragraph("DEMO / SYNTHETIC — учебный пример, не инструкция Jana Post.")
        document.add_heading("01. Учебный материал", level=1)
        document.add_paragraph(text)
        document.add_paragraph("Этот учебный пример нужен только для проверки поиска в назначенном курсе. "
                               "Он не описывает реальные рабочие правила или решения компании.")
        data = io.BytesIO()
        document.save(data)
        response = api.client.post("/api/knowledge-files", headers=api.headers("hr"),
            data={"title": name, "version": "test", "language": "ru",
                  "classification": "synthetic_demo", "access_scope": "all", "source_note": "test"},
            files={"file": (f"{name}.docx", data.getvalue())})
        assert response.status_code == 200, response.text
        return response.json()["id"]

    attached = upload("Көкшетау коды относится к сортировке посылок.", "attached")
    upload("Жарқырау коды относится к закрытому другому курсу.", "unattached")
    assert api.post("/api/courses/1/materials", {"file_id": attached}, "hr").status_code == 200
    seen = []
    monkeypatch.setattr(assistant, "generate_answer",
                        lambda question, passage, settings: seen.append(passage) or "Көкшетау коды описывает сортировку посылок.")
    result = api.post("/api/courses/1/ask", {"question": "Көкшетау коды"}).json()
    assert result["sources"][0]["file_id"] == attached
    assert "Жарқырау" not in seen[0]
    assert result["answer"].endswith(" [1]")
    assert api.post("/api/courses/1/ask", {"question": "Жарқырау коды"}).json()["sources"] == []
    assert len(seen) == 1
