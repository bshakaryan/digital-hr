"""Learning tracks on isolated PostgreSQL, with legacy rows and API retries."""
import io
from datetime import date

from alembic import command
from docx import Document
from sqlalchemy import text


def published_course(api, *, scope="all", self_enrollable=True):
    created = api.post("/api/courses", {"title": "Учебные статусы", "description": "Учебный пример",
        "objectives": "Различать события и предположения", "duration_minutes": 10,
        "self_enrollable": self_enrollable}, "hr")
    assert created.status_code == 201, created.text
    course_id = created.json()["id"]
    document = Document()
    document.add_paragraph("DEMO / SYNTHETIC — учебный пример, не инструкция Jana Post.")
    document.add_heading("01. Статусы", level=1)
    document.add_paragraph("Статус отправления отражает подтверждённое событие. Передача посылки между этапами фиксируется после фактического приёма. Предположение о местонахождении не является подтверждённым статусом. "
        f"Это учебный пример курса {course_id} с областью доступа {scope}.")
    stream = io.BytesIO()
    document.save(stream)
    uploaded = api.client.post("/api/knowledge-files", headers=api.headers("hr"),
        data={"title": f"Статусы {course_id}", "version": "test-1", "language": "ru",
              "classification": "synthetic_demo", "access_scope": scope},
        files={"file": (f"statuses-{course_id}.docx", stream.getvalue())})
    assert uploaded.status_code == 200, uploaded.text
    assert api.post(f"/api/courses/{course_id}/materials", {"file_id": uploaded.json()["id"]}, "hr").status_code == 200
    quiz = api.client.put(f"/api/courses/{course_id}/quiz", headers=api.headers("hr"),
        json={"question": "Что отражает статус?", "options": ["Событие", "Предположение"], "correct_index": 0})
    assert quiz.status_code == 200, quiz.text
    assert api.post(f"/api/courses/{course_id}/publish", {}, "hr").status_code == 200
    return course_id


def test_migration_preserves_populated_legacy_assignment(postgres_schema):
    command.upgrade(postgres_schema.config, "0012")
    with postgres_schema.engine.begin() as db:
        db.execute(text("""INSERT INTO employees
            (id,email,password_hash,full_name,role,company,position,department,hire_date,leave_balance)
            VALUES (1,'legacy@example.com','test','Legacy','employee','Demo','Role','Team','2025-01-01',12)"""))
        db.execute(text("""INSERT INTO courses
            (id,title,description,content,duration_minutes,question,options_json,correct_index)
            VALUES (1,'Legacy course','old','old',10,'Question','[\"yes\",\"no\"]',0)"""))
        db.execute(text("""INSERT INTO course_assignments
            (id,employee_id,course_id,progress,completed,deadline)
            VALUES (1,1,1,100,true,'2026-10-01')"""))
    command.upgrade(postgres_schema.config, "0013")
    with postgres_schema.engine.connect() as db:
        assert db.execute(text("SELECT learning_track,progress,completed,deadline FROM course_assignments WHERE id=1")).one() == (
            "legacy", 100, True, date(2026, 10, 1))
        assert db.scalar(text("SELECT self_enrollable FROM courses WHERE id=1")) is False
    command.upgrade(postgres_schema.config, "head")


def test_hr_assignment_and_audited_legacy_reclassification(api):
    course_id = published_course(api)
    employee_id = api.get("/api/me").json()["id"]
    catalog_item = next(item for item in api.get("/api/courses/catalog", "hr").json() if item["id"] == course_id)
    assert catalog_item["self_enrollable"] is True
    edited = api.patch(f"/api/courses/{course_id}", {"title": "Учебные статусы",
        "description": "Учебный пример", "objectives": "Различать события и предположения",
        "duration_minutes": 10}, "hr")
    assert edited.status_code == 200 and edited.json()["self_enrollable"] is True
    assert api.post(f"/api/courses/{course_id}/assign", {"employee_id": employee_id,
        "learning_track": "onboarding"}).status_code == 403
    assigned = api.post(f"/api/courses/{course_id}/assign", {"employee_id": employee_id,
        "learning_track": "onboarding"}, "hr")
    assert assigned.status_code == 200 and assigned.json()["assigned"] is True
    assert api.get(f"/api/courses/{course_id}").json()["learning_track"] == "onboarding"
    manager_id = next(person["id"] for person in api.get("/api/employees", "hr").json()
        if person["email"] == "manager@example.com")
    manager_assignment = api.post(f"/api/courses/{course_id}/assign", {"employee_id": manager_id,
        "learning_track": "development"}, "hr")
    assert manager_assignment.status_code == 200 and manager_assignment.json()["assigned"] is True
    assert api.get(f"/api/courses/{course_id}", "manager").json()["learning_track"] == "development"
    before = api.db.snapshot("course_assignments", "audit_events")
    assert api.post(f"/api/courses/{course_id}/assign", {"employee_id": employee_id,
        "learning_track": "development"}, "hr").json()["assigned"] is False
    assert api.post(f"/api/courses/{course_id}/enroll", {}).json()["learning_track"] == "onboarding"
    assert api.db.snapshot("course_assignments", "audit_events") == before
    path = f"/api/courses/1/assignments/{employee_id}/track"
    assert api.patch(path, {"learning_track": "development"}).status_code == 403
    assert api.patch(path, {"learning_track": "legacy"}, "hr").status_code == 422
    original = api.get("/api/courses/1").json()
    changed = api.patch(path, {"learning_track": "development"}, "hr")
    assert changed.status_code == 200 and changed.json()["changed"] is True
    audit = api.db.snapshot("audit_events")
    assert api.patch(path, {"learning_track": "development"}, "hr").json()["changed"] is False
    assert api.patch(path, {"learning_track": "onboarding"}, "hr").status_code == 409
    assert api.db.snapshot("audit_events") == audit
    current = api.get("/api/courses/1").json()
    assert current["learning_track"] == "development"
    assert (current["progress"], current["completed"], current["deadline"]) == (
        original["progress"], original["completed"], original["deadline"])


def test_self_enrollment_access_idempotence_and_xp(api):
    course_id = published_course(api)
    assert api.get(f"/api/courses/{course_id}").status_code == 404
    before_xp = api.get("/api/courses/rewards").json()["xp"]
    enrolled = api.post(f"/api/courses/{course_id}/enroll", {})
    assert enrolled.status_code == 200 and enrolled.json()["learning_track"] == "development"
    snapshot = api.db.snapshot("course_assignments", "audit_events")
    assert api.post(f"/api/courses/{course_id}/enroll", {}).json()["assigned"] is False
    assert api.db.snapshot("course_assignments", "audit_events") == snapshot
    assert api.get(f"/api/courses/{course_id}").json()["learning_track"] == "development"
    assert api.get("/api/courses/rewards").json()["xp"] == before_xp
    assert api.post(f"/api/courses/{course_id}/progress", {"progress": 70}).status_code == 200
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 0}).status_code == 200
    assert api.get("/api/courses/rewards").json()["xp"] == before_xp + 100
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 0}).status_code == 200
    assert api.get("/api/courses/rewards").json()["xp"] == before_xp + 100


def test_self_enrollment_requires_published_accessible_opt_in(api):
    draft = api.post("/api/courses", {"title": "Черновик", "description": "Описание", "objectives": "Цель",
        "duration_minutes": 10, "self_enrollable": True}, "hr").json()
    assert api.post(f"/api/courses/{draft['id']}/enroll", {}).status_code == 404
    closed = published_course(api, self_enrollable=False)
    assert api.post(f"/api/courses/{closed}/enroll", {}).status_code == 409
    restricted = published_course(api, scope="manager")
    assert api.post(f"/api/courses/{restricted}/enroll", {}).status_code == 404
    assert api.post(f"/api/courses/{restricted}/enroll", {}, "manager").status_code == 200
    assert api.get(f"/api/courses/{restricted}").status_code == 404
    assert api.get(f"/api/courses/{restricted}", "manager").status_code == 200
