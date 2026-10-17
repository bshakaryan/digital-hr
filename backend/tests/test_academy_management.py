"""HR authoring and employee completion through the public Academy API."""
import io

from docx import Document


def document_bytes(extra=""):
    document = Document()
    document.add_paragraph("DEMO / SYNTHETIC — учебный пример, не инструкция Jana Post.")
    document.add_heading("01. События отслеживания", level=1)
    document.add_paragraph("Статус отправления отражает подтверждённое событие. Передача посылки между этапами фиксируется после фактического приёма. Предположение о местонахождении не является подтверждённым статусом.")
    if extra:
        document.add_paragraph(extra)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def uploaded_file(api, scope="all"):
    response = api.client.post("/api/knowledge-files", headers=api.headers("hr"),
        data={"title": "Учебные события", "version": "test-1", "language": "ru",
              "classification": "synthetic_demo", "access_scope": scope},
        files={"file": ("events.docx", document_bytes("Только HR." if scope == "hr_admin" else ""))})
    assert response.status_code == 200, response.text
    return response.json()["id"]


def draft(api):
    response = api.post("/api/courses", {"title": "Статусы отправлений", "description": "Учебный курс об отслеживании",
        "objectives": "Отличать подтверждённый статус от предположения", "duration_minutes": 12}, "hr")
    assert response.status_code == 201, response.text
    return response.json()


def quiz(api, course_id):
    response = api.client.put(f"/api/courses/{course_id}/quiz", headers=api.headers("hr"), json={
        "question": "Что отражает статус отправления?",
        "options": ["Подтверждённое событие", "Предположение о местонахождении"], "correct_index": 0})
    assert response.status_code == 200, response.text


def test_hr_builds_publishes_assigns_and_employee_completes_document_course(api):
    before = api.db.snapshot("courses", "course_assignments", "notifications", "audit_events")
    item = draft(api)
    course_id = item["id"]
    assert item["status"] == "draft"
    assert all(course["id"] != course_id for course in api.get("/api/courses").json())
    assert api.get(f"/api/courses/{course_id}").status_code == 404
    assert api.post(f"/api/courses/{course_id}/publish", {}, "hr").status_code == 409
    edited = api.patch(f"/api/courses/{course_id}", {"title": "События отправлений", "description": "Факты и статусы",
        "objectives": "Проверить факт передачи посылки", "duration_minutes": 9}, "hr")
    assert edited.status_code == 200 and edited.json()["title"] == "События отправлений"
    assert api.post(f"/api/courses/{course_id}/assign", {"employee_id": 1}, "hr").status_code == 409
    file_id = uploaded_file(api)
    assert api.post(f"/api/courses/{course_id}/materials", {"file_id": file_id}, "hr").status_code == 200
    quiz(api, course_id)
    assert api.post(f"/api/courses/{course_id}/publish", {}, "hr").json()["status"] == "published"
    published = api.db.snapshot("audit_events")
    assert api.post(f"/api/courses/{course_id}/publish", {}, "hr").status_code == 200
    assert api.db.snapshot("audit_events") == published
    assigned = api.post(f"/api/courses/{course_id}/assign", {"employee_id": 1}, "hr")
    assert assigned.status_code == 200 and assigned.json()["assigned"] is True
    after_assignment = api.db.snapshot("course_assignments", "notifications", "audit_events")
    assert api.post(f"/api/courses/{course_id}/assign", {"employee_id": 1}, "hr").json()["assigned"] is False
    assert api.db.snapshot("course_assignments", "notifications", "audit_events") == after_assignment
    manager_id = next(person["id"] for person in api.get("/api/employees", "hr").json() if person["email"] == "manager@example.com")
    second_assignment = api.post(f"/api/courses/{course_id}/assign", {"employee_id": manager_id}, "hr")
    assert second_assignment.status_code == 200 and second_assignment.json()["assigned"] is True
    assert api.get(f"/api/courses/{course_id}", "manager").status_code == 200
    both_assigned = api.db.snapshot("course_assignments", "notifications", "audit_events")
    assert api.post(f"/api/courses/{course_id}/assign", {"employee_id": manager_id}, "hr").json()["assigned"] is False
    assert api.db.snapshot("course_assignments", "notifications", "audit_events") == both_assigned
    course = api.get(f"/api/courses/{course_id}").json()
    assert course["progress"] == 0 and course["objectives"] == "Проверить факт передачи посылки"
    assert [file["id"] for file in api.get(f"/api/courses/{course_id}/materials").json()] == [file_id]
    answer = api.post(f"/api/courses/{course_id}/ask", {"question": "Что отражает статус отправления?"}).json()
    assert answer["sources"] and answer["sources"][0]["file_id"] == file_id
    assert answer["sources"][0]["section"] == "01. События отслеживания"
    assert "подтверждённое событие" in answer["answer"].lower()
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 0}).status_code == 409
    assert api.post(f"/api/courses/{course_id}/progress", {"progress": 70}).json()["progress"] == 70
    assert api.get(f"/api/courses/{course_id}").json()["progress"] == 70
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 1}).status_code == 422
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 0}).json()["completed"] is True
    assert api.get(f"/api/courses/{course_id}").json()["progress"] == 100
    assert api.get("/api/courses/rewards").json()["xp"] == 100
    completed = api.db.snapshot("course_assignments", "notifications", "audit_events")
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 0}).status_code == 200
    assert api.db.snapshot("course_assignments", "notifications", "audit_events") == completed
    assert api.get("/api/courses/rewards").json()["xp"] == 100
    assert len(api.db.snapshot("courses")["courses"]) == len(before["courses"]) + 1


def test_authoring_permissions_validation_and_published_quiz_lock(api):
    body = {"title": "Тест", "description": "Описание", "objectives": "Цели", "duration_minutes": 10}
    assert api.post("/api/courses", body).status_code == 403
    assert api.get("/api/courses/catalog").status_code == 403
    item = draft(api)
    course_id = item["id"]
    assert api.patch(f"/api/courses/{course_id}", body).status_code == 403
    assert api.client.put(f"/api/courses/{course_id}/quiz", headers=api.headers(), json={"question": "Вопрос?", "options": ["Да", "Нет"], "correct_index": 0}).status_code == 403
    assert api.post(f"/api/courses/{course_id}/publish", {}).status_code == 403
    assert api.post(f"/api/courses/{course_id}/assign", {"employee_id": 1}).status_code == 403
    assert api.client.put(f"/api/courses/{course_id}/quiz", headers=api.headers("hr"), json={"question": "Вопрос?", "options": ["Да", "Да"], "correct_index": 0}).status_code == 422
    quiz(api, course_id)
    restricted_file = uploaded_file(api, "hr_admin")
    assert api.post(f"/api/courses/{course_id}/materials", {"file_id": restricted_file}, "hr").status_code == 200
    assert api.post(f"/api/courses/{course_id}/publish", {}, "hr").status_code == 409
    public_file = uploaded_file(api)
    assert api.post(f"/api/courses/{course_id}/materials", {"file_id": public_file}, "hr").status_code == 200
    assert api.post(f"/api/courses/{course_id}/publish", {}, "hr").status_code == 200
    assert api.client.put(f"/api/courses/{course_id}/quiz", headers=api.headers("hr"), json={"question": "Другой вопрос?", "options": ["Да", "Нет"], "correct_index": 0}).status_code == 409
    assert api.post(f"/api/courses/{course_id}/assign", {"employee_id": 99999}, "hr").status_code == 404
    assert api.get(f"/api/courses/{course_id}", "outsider").status_code == 404
