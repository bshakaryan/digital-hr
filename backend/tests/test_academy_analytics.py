"""Deadline, analytics, discovery and learning-profile authorization."""
import io

from docx import Document
from sqlalchemy import select


def prepared_course(api):
    document = Document()
    document.add_paragraph("DEMO / SYNTHETIC — учебный пример, не инструкция Jana Post.")
    document.add_heading("01. События", level=1)
    document.add_paragraph("Статус отражает подтверждённое событие, а не предположение. "
        "Физический приём посылки подтверждают отдельной записью с временем и местом. "
        "Если подтверждения нет, сотрудник проверяет источник данных и не утверждает, что отправление уже доставлено.")
    stream = io.BytesIO()
    document.save(stream)
    uploaded = api.client.post("/api/knowledge-files", headers=api.headers("hr"),
        data={"title": "Учебный документ", "version": "1.0", "language": "ru",
              "classification": "synthetic_demo", "access_scope": "all"},
        files={"file": ("events.docx", stream.getvalue())})
    assert uploaded.status_code == 200, uploaded.text
    created = api.post("/api/courses", {"title": "Учебные статусы", "description": "Демо о событиях",
        "objectives": "Отличать событие от предположения", "duration_minutes": 12, "category": "logistics"}, "hr")
    assert created.status_code == 201, created.text
    course_id = created.json()["id"]
    assert api.post(f"/api/courses/{course_id}/materials", {"file_id": uploaded.json()["id"]}, "hr").status_code == 200
    quiz = api.client.put(f"/api/courses/{course_id}/quiz", headers=api.headers("hr"), json={
        "question": "Что отражает статус?", "options": ["Подтверждённое событие", "Предположение"], "correct_index": 0})
    assert quiz.status_code == 200
    assert api.post(f"/api/courses/{course_id}/publish", {}, "hr").status_code == 200
    return course_id


def test_deadline_metrics_profiles_discovery_and_permissions(api):
    course_id = prepared_course(api)
    discovered = api.get("/api/courses/discover").json()
    assert any(item["id"] == course_id and item["category"] == "logistics" for item in discovered)
    assert all("correct_index" not in item for item in discovered)
    assigned = api.post(f"/api/courses/{course_id}/assign", {"employee_id": 1, "deadline": "2026-10-16"}, "hr")
    assert assigned.status_code == 200 and assigned.json()["assigned"]
    repeated = api.post(f"/api/courses/{course_id}/assign", {"employee_id": 1, "deadline": "2026-10-30"}, "hr")
    assert repeated.status_code == 200 and not repeated.json()["assigned"]
    assert api.get(f"/api/courses/{course_id}").json()["deadline"] == "2026-10-16"
    assert api.get("/api/courses/analytics").status_code == 403
    assert api.get("/api/courses/analytics", "manager").status_code == 403
    assert api.get("/api/courses/analytics/employees").status_code == 403
    assert api.get("/api/courses/analytics/employees/4").status_code == 403
    assert api.get("/api/courses/analytics/employees/4", "manager").status_code == 404
    assert api.get("/api/courses/analytics/employees/1", "manager").status_code == 200
    assert api.get("/api/courses/analytics/employees/4", "hr").status_code == 200
    directory = api.get("/api/courses/analytics/employees", "manager").json()
    assert [employee["employee_id"] for employee in directory] == [1]
    before = api.get("/api/courses/analytics", "hr").json()
    assert before["total_demo_employees"] == 4
    assert before["total_assignments"] == 5
    assert before["overdue_assignments"] == 1
    assert before["completion_rate"] == 0
    assert sum(before["status_counts"].values()) == before["total_assignments"]
    assert api.post(f"/api/courses/{course_id}/progress", {"progress": 70}).status_code == 200
    profile = api.get("/api/courses/analytics/employees/1", "hr").json()
    item = next(item for item in profile["assignments"] if item["course_id"] == course_id)
    assert item["status"] == "in_progress" and item["overdue"] is True
    assert item["last_activity_at"] is not None and item["completed_at"] is None
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 1}).status_code == 422
    assert api.post(f"/api/courses/{course_id}/complete", {"answer_index": 0}).status_code == 200
    after = api.get("/api/courses/analytics", "hr").json()
    assert after["completion_rate"] == 20.0
    assert after["overdue_assignments"] == 0
    assert after["status_counts"]["completed"] == 1
    profile = api.get("/api/courses/analytics/employees/1", "hr").json()
    assert profile["xp"] == api.get("/api/courses/rewards").json()["xp"] == 100
    assert profile["completion_rate"] == 50.0
    item = next(item for item in profile["assignments"] if item["course_id"] == course_id)
    assert item["status"] == "completed" and item["overdue"] is False
    assert item["completed_at"] is not None


def test_demo_provision_is_repeatable_and_preserves_existing_progress(api, monkeypatch):
    from app.models import Course, CourseAssignment, KnowledgeDocument, KnowledgeFile
    from app import provision_academy_demo as demo

    COURSES, ONBOARDING, SELF_ENROLLABLE = demo.COURSES, demo.ONBOARDING, demo.SELF_ENROLLABLE

    with api.db.session() as db:
        hr_id = 3
        files = {}
        for index, code in enumerate(sorted({code for spec in COURSES for code in spec.files})):
            file = KnowledgeFile(sha256=f"{index + 1:064x}", filename=f"{code}.docx", title=code,
                document_code=code, version="test", language="ru", classification="synthetic_demo",
                source_note="DEMO / SYNTHETIC", access_scope="all", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                size_bytes=1, original_bytes=b"x", uploaded_by=hr_id)
            db.add(file)
            db.flush()
            files[code] = file
        for spec in COURSES:
            db.add(KnowledgeDocument(title=spec.title, section=spec.section, content=spec.evidence,
                keywords="", access_scope="all", file_id=files[spec.files[0]].id, chunk_index=0))
        old = db.scalar(select(CourseAssignment).where(CourseAssignment.employee_id == 1, CourseAssignment.course_id == 1))
        old.progress = 45
        db.commit()
    # Reproduce the reported baseline before enabling onboarding and discovery.
    with monkeypatch.context() as prior:
        prior.setattr(demo, "ONBOARDING", {})
        prior.setattr(demo, "SELF_ENROLLABLE", frozenset())
        baseline = demo.provision()
    assert baseline["created_courses"] == 12 and baseline["new_assignments"] == 16
    assert baseline["completed"] == 4
    old_before = api.db.snapshot("course_assignments")
    assert len(old_before["course_assignments"]) == 20
    with api.db.session() as db:
        assert all(item.learning_track == "legacy" for item in db.scalars(select(CourseAssignment)).all())
        assert all(not item.self_enrollable for item in db.scalars(select(Course).where(Course.demo_code.is_not(None))).all())
    assert api.get("/api/courses/rewards").json()["xp"] == 200

    first = demo.provision()
    assert first["created_courses"] == 0 and first["new_assignments"] == 4
    assert first["new_onboarding_assignments"] == 4 and first["self_enrollable_enabled"] == 3
    assert first["completed"] == 1
    assigned = api.get("/api/courses").json()
    onboarding = [course for course in assigned if course["learning_track"] == "onboarding"]
    assert len(onboarding) == 4
    assert {course["title"] for course in onboarding} == {
        f"Демо · {spec.title}" for spec in COURSES if spec.code in {code for code, _, _ in ONBOARDING["employee@example.com"]}}
    assert sum(course["completed"] for course in onboarding) == 1
    assert sorted(course["progress"] for course in onboarding) == [0, 0, 70, 100]
    assert not any(course["learning_track"] == "onboarding" for course in api.get("/api/courses", "outsider").json())
    assert api.get("/api/courses/rewards").json()["xp"] == 300
    assert api.db.snapshot("course_assignments")["course_assignments"][:20] == old_before["course_assignments"]
    catalog = api.get("/api/courses/discover").json()
    assert sum(course["self_enrollable"] for course in catalog if course["is_demo"]) == 3
    with api.db.session() as db:
        assert sum(1 for code in SELF_ENROLLABLE if db.scalar(select(Course).where(Course.demo_code == code)).self_enrollable) == 3
    assert api.get("/api/courses/catalog").status_code == 403
    assert api.get("/api/courses/catalog", "hr").status_code == 200
    assert api.get("/api/courses/analytics", "hr").status_code == 200
    assert api.post(f"/api/courses/{onboarding[0]['id']}/assign", {"employee_id": 4}).status_code == 403
    snapshot = api.db.snapshot("courses", "course_assignments", "audit_events", "notifications")
    second = demo.provision()
    assert second["created_courses"] == 0 and second["new_assignments"] == 0
    assert second["new_onboarding_assignments"] == 0 and second["self_enrollable_enabled"] == 0
    assert api.db.snapshot("courses", "course_assignments", "audit_events", "notifications") == snapshot
    assert api.get("/api/courses/rewards").json()["xp"] == 300
    # A previously assigned legacy course cannot gain an onboarding track by
    # appearing in a future demo curriculum.
    with monkeypatch.context() as conflict:
        conflict.setattr(demo, "ONBOARDING", {"employee@example.com": (("DEMO-ONB-01", "completed", 7),)})
        assert demo.provision()["new_onboarding_assignments"] == 0
    assert api.db.snapshot("courses", "course_assignments", "audit_events", "notifications") == snapshot
    with api.db.session() as db:
        old = db.scalar(select(CourseAssignment).where(CourseAssignment.employee_id == 1, CourseAssignment.course_id == 1))
        assert old.progress == 45 and not old.completed and old.learning_track == "legacy"
        assert api.db.snapshot("course_assignments")["course_assignments"][:len(old_before["course_assignments"])] == old_before["course_assignments"]
