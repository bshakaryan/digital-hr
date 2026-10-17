"""Shared HR upload and initial-corpus ingestion service; originals stay byte-for-byte intact."""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models import KnowledgeDocument, KnowledgeFile
from app.modules.audit import record

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_CHUNK_CHARS = 1100
CLASSIFICATIONS = {"synthetic_demo", "external_legal_snapshot", "external_reference"}
SCOPES = {"all", "employee", "manager", "hr_admin"}
HEADING = re.compile(r"^(?:\d{2}\.|Статья\s+\d+[.-]|Глава\s+\d+[.-]|РАЗДЕЛ\s+\d+[.-])\s*.+", re.IGNORECASE)


@dataclass(frozen=True)
class Chunk:
    section: str
    page_number: int | None
    content: str


@dataclass(frozen=True)
class UploadDetails:
    title: str
    version: str
    language: str
    classification: str
    access_scope: str
    source_note: str = ""
    document_code: str | None = None


def _safe_filename(filename: str) -> str:
    name = PurePath(filename.replace("\\", "/")).name.strip()
    if not name or len(name) > 255 or name.startswith("."):
        raise ValueError("Некорректное имя файла")
    if name.lower().split(".")[-1] not in {"pdf", "docx"}:
        raise ValueError("Поддерживаются только PDF и DOCX")
    return name


def _split_text(text: str, limit: int = MAX_CHUNK_CHARS) -> list[str]:
    normalized = re.sub(r"[ \t]+", " ", text).strip()
    if not normalized:
        return []
    if len(normalized) <= limit:
        return [normalized]
    pieces = []
    while normalized:
        if len(normalized) <= limit:
            pieces.append(normalized)
            break
        cut = normalized.rfind(" ", 0, limit + 1)
        if cut < limit // 2:
            cut = limit
        pieces.append(normalized[:cut].strip())
        normalized = normalized[cut:].strip()
    return pieces


def _group_lines(lines: list[str], section: str, page_number: int | None) -> list[Chunk]:
    result, buffer = [], []
    size = 0
    for line in lines:
        for part in _split_text(line):
            if buffer and size + len(part) + 1 > MAX_CHUNK_CHARS:
                result.append(Chunk(section, page_number, "\n".join(buffer)))
                buffer, size = [], 0
            buffer.append(part)
            size += len(part) + 1
    if buffer:
        result.append(Chunk(section, page_number, "\n".join(buffer)))
    return result


def extract_pdf(data: bytes) -> list[Chunk]:
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ValueError("Зашифрованный PDF не поддерживается")
        chunks = []
        heading = "Общий текст"
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise ValueError(f"PDF, страница {number}: текст не извлечён; OCR не подключён")
            lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
            buffer = []
            for line in lines:
                if HEADING.match(line) and len(line) <= 120:
                    chunks.extend(_group_lines(buffer, heading, number))
                    heading = line
                    buffer = []
                else:
                    buffer.append(line)
            chunks.extend(_group_lines(buffer, heading, number))
        return chunks
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("PDF не удалось прочитать") from exc


def extract_docx(data: bytes) -> list[Chunk]:
    try:
        document = Document(io.BytesIO(data))
        paragraphs = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        chunks = []
        heading, buffer = "Общий текст", []
        for paragraph in paragraphs:
            line = re.sub(r"\s+", " ", paragraph.text).strip()
            if not line:
                continue
            is_heading = paragraph.style.name.startswith("Heading") or bool(HEADING.match(line) and len(line) <= 120)
            if is_heading:
                chunks.extend(_group_lines(buffer, heading, None))
                heading, buffer = line[:120], []
            else:
                buffer.append(line)
        chunks.extend(_group_lines(buffer, heading, None))
        return chunks
    except Exception as exc:
        raise ValueError("DOCX не удалось прочитать") from exc


def ingest_file(db: Session, *, data: bytes, filename: str, details: UploadDetails, actor_id: int) -> tuple[KnowledgeFile, bool]:
    name = _safe_filename(filename)
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("Файл пустой или превышает 10 МБ")
    if not 1 <= len(details.title.strip()) <= 200 or not 1 <= len(details.version.strip()) <= 80:
        raise ValueError("Укажите название и версию документа")
    if details.classification not in CLASSIFICATIONS or details.access_scope not in SCOPES:
        raise ValueError("Некорректная классификация или область доступа")
    if details.language not in {"ru", "kk", "en"}:
        raise ValueError("Язык должен быть ru, kk или en")
    if details.document_code and len(details.document_code) > 40:
        raise ValueError("Слишком длинный идентификатор документа")
    if len(details.source_note) > 2000:
        raise ValueError("Слишком длинное примечание об источнике")
    digest = hashlib.sha256(data).hexdigest()
    existing = db.scalar(select(KnowledgeFile).where(KnowledgeFile.sha256 == digest))
    if existing is not None:
        return existing, False
    is_pdf = name.lower().endswith(".pdf")
    if is_pdf and not data.startswith(b"%PDF-"):
        raise ValueError("Расширение PDF не соответствует содержимому")
    if not is_pdf and not data.startswith(b"PK"):
        raise ValueError("Расширение DOCX не соответствует содержимому")
    chunks = extract_pdf(data) if is_pdf else extract_docx(data)
    if not chunks or sum(len(chunk.content) for chunk in chunks) < 100:
        raise ValueError("В документе не найден пригодный для поиска текст")
    if details.classification == "synthetic_demo":
        preview = " ".join(chunk.content for chunk in chunks[:5]).lower()
        if not any(marker in preview for marker in ("demo", "synthetic", "демонстрацион")):
            raise ValueError("Синтетический документ должен явно содержать DEMO / SYNTHETIC")
    item = KnowledgeFile(
        sha256=digest, filename=name, title=details.title.strip(),
        document_code=details.document_code.strip() if details.document_code else None,
        version=details.version.strip(), language=details.language,
        classification=details.classification, source_note=details.source_note.strip(),
        access_scope=details.access_scope,
        media_type="application/pdf" if is_pdf else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=len(data), original_bytes=data, uploaded_by=actor_id,
    )
    try:
        db.add(item)
        db.flush()
        for index, chunk in enumerate(chunks, 1):
            db.add(KnowledgeDocument(
                title=item.title, section=chunk.section, content=chunk.content,
                keywords=" ".join(filter(None, (item.document_code, chunk.section))),
                access_scope=item.access_scope, file_id=item.id,
                page_number=chunk.page_number, chunk_index=index,
            ))
        record(db, actor_id, "knowledge.imported", "knowledge_file", item.id)
        db.commit()
    except IntegrityError:
        db.rollback()
        existing = db.scalar(select(KnowledgeFile).where(KnowledgeFile.sha256 == digest))
        if existing is not None:
            return existing, False
        raise
    db.refresh(item)
    return item, True


def chunk_count(db: Session, file_id: int) -> int:
    return db.scalar(select(func.count()).select_from(KnowledgeDocument).where(KnowledgeDocument.file_id == file_id)) or 0
