"""Read-only inventory of the provided knowledge_base_docs corpus."""
from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
import fitz

root = Path(__file__).resolve().parents[1] / "knowledge_base_docs"
for path in sorted(root.iterdir()):
    if not path.is_file():
        continue
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix.lower() == ".pdf":
        document = fitz.open(path)
        pages = len(document)
        text = "\n".join(page.get_text() for page in document)
        kind = f"PDF, {pages} pages"
    else:
        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        kind = f"DOCX, {len(document.paragraphs)} paragraphs"
    print(f"FILE {path.name}\n  {kind}; chars={len(text)}; sha256={digest}")
    print(f"  Cyrillic={any('А' <= c <= 'я' for c in text)}; Kazakh={any(c in 'әіңғұқөүһӘІҢҒҰҚӨҮҺ' for c in text)}")
    print("  START " + text[:420].replace("\n", " | "))
